from io import BytesIO
from docx import Document
from openpyxl import Workbook
import fitz
import pytest
from orgx.ingest import ingest
from orgx.extract import extract
from orgx.engine import audit


def test_docx_exact_text_table_order_and_stable_ids():
    d = Document()
    d.add_paragraph("3. Структура")
    d.add_paragraph("3.4. БВА состоит из следующих структурных подразделений:")
    d.add_paragraph("а. Департамент тестирования (ДТ).")
    d.add_paragraph("")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "  точный  текст\nв ячейке"
    t.cell(0, 1).text = "вторая"
    d.add_paragraph("после таблицы")
    f = BytesIO()
    d.save(f)
    doc = ingest(f.getvalue(), "input.docx", "before")
    assert doc.spans[3].exact_text == "  точный  текст\nв ячейке"
    assert doc.spans[3].paragraph_id == "t0r0c0p0"
    assert doc.spans[-1].paragraph_id == "p0004"
    assert doc == ingest(f.getvalue(), "input.docx", "before")
    assert doc.spans[2].clause == "3.4/а"


def test_xlsx_preserves_sheet_cell_locator_and_formula():
    book = Workbook()
    book.active.title = "Owners"
    book.active["A1"] = "5.3. Подразделение A:"
    book.active["A2"] = "5.3.1. Проверяет отчётность"
    book.active["B2"] = "=1+1"
    f = BytesIO()
    book.save(f)
    doc = ingest(f.getvalue(), "input.xlsx", "after")
    assert doc.spans[-1].exact_text == "=1+1"
    assert doc.spans[-1].locator == "Owners!B2"


def test_pdf_and_empty_scan():
    d = fitz.open()
    page = d.new_page()
    page.insert_text((50, 50), "5.3.1. Reviews controls")
    doc = ingest(d.tobytes(), "input.pdf", "before")
    assert "Reviews controls" in doc.spans[0].exact_text
    assert doc.spans[0].locator.startswith("Страница 1")
    empty = fitz.open()
    empty.new_page()
    with pytest.raises(ValueError, match="OCR"):
        ingest(empty.tobytes(), "scan.pdf", "before")


@pytest.mark.parametrize(
    "name,data", [("bad.docx", b"not zip"), ("bad.exe", b"file"), ("empty.docx", b"")]
)
def test_invalid_upload_rejected(name, data):
    with pytest.raises(ValueError):
        ingest(data, name, "before")


def test_unknown_structure_requires_review():
    d = Document()
    d.add_paragraph("Произвольный документ без обязанностей")
    f = BytesIO()
    d.save(f)
    r = audit(
        extract(
            [
                ingest(f.getvalue(), "old.docx", "before"),
                ingest(f.getvalue(), "new.docx", "after"),
            ]
        )
    )
    assert r.findings[0].type == "REQUIRES_HUMAN_REVIEW"


def test_unknown_heading_does_not_inherit_previous_owner():
    d = Document()
    for text in [
        "5. Права и обязанности",
        "5.3. Подразделение A:",
        "5.3.1. Проверяет отчётность.",
        "5.4. Особый порядок:",
        "5.4.1. Назначает ответственного.",
    ]:
        d.add_paragraph(text)
    f = BytesIO()
    d.save(f)
    ir = extract([ingest(f.getvalue(), "input.docx", "before")])
    assert len(ir.claims) == 2 and "Проверяет" in ir.claims[0].text
    assert ir.claims[1].extraction == "unresolved"
    assert (
        next(u for u in ir.units if u.id == ir.claims[1].unit_id).kind == "unresolved"
    )
