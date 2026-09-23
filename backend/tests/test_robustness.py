"""Control documents authored independently of the demo pair and its truth labels.

These are engineering checks, not independent expert acceptance. Every scenario
goes through DOCX ingestion and extraction rather than supplying a fabricated IR.
"""

from io import BytesIO
import re

from docx import Document
import pytest

from orgx.engine import audit, validate_record
from orgx.extract import extract
from orgx.ingest import ingest


def document(lines, version):
    word = Document()
    for line in lines:
        word.add_paragraph(line)
    data = BytesIO()
    word.save(data)
    return ingest(data.getvalue(), version + ".docx", version)


def rewrite_original(original, transform):
    from conftest import ROOT

    word = Document(ROOT / "work/hackalem/kazakhtelecom" / original.name)
    for p in word.paragraphs:
        updated = transform(p.text)
        if updated != p.text:
            p.text = updated
    data = BytesIO()
    word.save(data)
    return ingest(data.getvalue(), original.name, original.version)


DUTY = "Проверяет полноту данных реестра оборудования за отчётный месяц."


def control(owner="Директор ДА:", duty=DUTY, context=None):
    lines = [
        "3. Структура",
        "3.9. Служба состоит из следующих структурных подразделений:",
        "а. Департамент анализа (ДА).",
        "б. Департамент проверки (ДП).",
        "5. Права и обязанности",
        "5.8. " + owner,
    ]
    if context:
        lines.extend(["5.8.4. " + context, "а. " + duty])
    else:
        lines.append("5.8.4. " + duty)
    return lines


@pytest.mark.parametrize(
    "before,after,expected",
    [
        pytest.param(control(), control(), "PRESERVED", id="same-owner-and-duty"),
        pytest.param(
            control(),
            control(duty="Оценивает полноту сведений об оборудовании ежемесячно."),
            "REQUIRES_HUMAN_REVIEW",
            id="paraphrase-is-a-candidate",
        ),
        pytest.param(
            control(context="В отношении всех филиалов:"),
            control(context="В отношении северного филиала:"),
            "REQUIRES_HUMAN_REVIEW",
            id="narrower-parent-context",
        ),
        pytest.param(
            control(owner="Директор ДА имеет право в отношении всех филиалов:"),
            control(owner="Директор ДА имеет право в отношении северного филиала:"),
            "REQUIRES_HUMAN_REVIEW",
            id="narrower-owner-context",
        ),
        pytest.param(
            control(owner="Директор ДА имеет право:"),
            control(owner="Директор ДА не имеет права:"),
            "REQUIRES_HUMAN_REVIEW",
            id="right-to-prohibition",
        ),
        pytest.param(
            control(),
            control(owner="Особый порядок:"),
            "REQUIRES_HUMAN_REVIEW",
            id="missing-owner",
        ),
        pytest.param(
            control(),
            control(owner="Директор ДП:"),
            "REQUIRES_HUMAN_REVIEW",
            id="new-owner-is-not-proven-transfer",
        ),
        pytest.param(
            control(),
            control(duty="Ведёт переписку по вопросам обучения персонала."),
            "REQUIRES_HUMAN_REVIEW",
            id="no-candidate-is-not-proven-loss",
        ),
        pytest.param(
            control(),
            control(
                duty="Не проверяет полноту данных реестра оборудования за отчётный месяц."
            ),
            "REQUIRES_HUMAN_REVIEW",
            id="negated-duty",
        ),
    ],
)
def test_independent_control_documents(before, after, expected):
    result = audit(extract([document(before, "before"), document(after, "after")]))
    target = next(
        c for c in result.ir.claims if c.version == "before" and c.text == DUTY
    )
    findings = [f for f in result.findings if f.before_claim_id == target.id]
    assert len(findings) == 1
    assert findings[0].type == expected
    assert findings[0].review_status == "PENDING"
    assert findings[0].source_span_ids
    search = next(s for s in result.search_results if s.before_claim == target.id)
    assert set(search.searched_span_ids) == {s.id for s in result.ir.documents[1].spans}
    validate_record(result)


def test_independent_duplicate_control():
    after = control() + ["5.9. Директор ДП:", "5.9.4. " + DUTY]
    result = audit(extract([document(control(), "before"), document(after, "after")]))
    duplicate = [f for f in result.findings if f.type == "POTENTIAL_DUPLICATE"]
    assert len(duplicate) == 1
    assert len(duplicate[0].after_claim_ids) == 2
    validate_record(result)


def test_unknown_heading_at_former_chief_auditor_number_does_not_invent_owner():
    lines = ["5. Права и обязанности", "5.1. Особый порядок:", "5.1.4. " + DUTY]
    result = extract([document(lines, "before")])
    assert result.claims[0].extraction == "unresolved"
    assert all(u.name != "Главный аудитор" for u in result.units)


@pytest.mark.parametrize(
    "heading,expected",
    [
        ("Директор ДА имеет право:", "right"),
        ("Директор ДА не имеет права:", "prohibition"),
        ("Директор ДА обязан:", "duty"),
    ],
)
def test_modality_is_separate_from_owner_identity(heading, expected):
    ir = extract([document(control(owner=heading), "before")])
    claim = ir.claims[0]
    assert claim.authority == expected
    assert next(u for u in ir.units if u.id == claim.unit_id).name == "Директор ДА"


def test_scope_contains_source_qualifier_not_owner_name():
    ir = extract(
        [
            document(
                control(owner="Директор ДА имеет право в отношении северного филиала:"),
                "before",
            )
        ]
    )
    assert "северного филиала" in ir.claims[0].scope
    assert "Директор ДА" not in ir.claims[0].scope
    assert extract([document(control(), "after")]).claims[0].scope == ""


def test_changed_leading_quantity_is_not_stripped_twice():
    before = control(duty="10 проверок реестра оборудования проводит ежегодно.")
    after = control(duty="20 проверок реестра оборудования проводит ежегодно.")
    result = audit(extract([document(before, "before"), document(after, "after")]))
    assert [f.type for f in result.findings] == ["REQUIRES_HUMAN_REVIEW"]
    assert result.search_results[0].exact_match_count == 0


def test_whitespace_case_and_letter_label_do_not_change_a_duty():
    before = control(context="В отношении всех филиалов:")
    after = ["  " + re.sub(r" ", "  ", line.upper()) for line in before]
    result = audit(extract([document(before, "before"), document(after, "after")]))
    target = [
        f
        for f in result.findings
        if f.before_claim_id
        and next(c for c in result.ir.claims if c.id == f.before_claim_id).text == DUTY
    ]
    assert len(target) == 1 and target[0].type == "PRESERVED"


def renumber(lines, offset):
    def replace(match):
        return ".".join(str(int(n) + offset) for n in match[1].split(".")) + match[2]

    return [re.sub(r"^(\d+(?:\.\d+)*)(\.\s|\s)", replace, line) for line in lines]


def signature(record):
    # IDs/hashes and clause labels must change with source bytes; conclusions and
    # physical source positions must not change under a pure relabelling.
    spans = {s.id: s for d in record.ir.documents for s in d.spans}
    claims = {c.id: c for c in record.ir.claims}

    def position(cid):
        return spans[claims[cid].source_span_ids[0]].paragraph_id if cid else None

    return sorted(
        (
            f.type,
            f.rule_id,
            position(f.before_claim_id) or "",
            tuple(sorted(position(c) for c in f.after_claim_ids)),
        )
        for f in record.findings
    )


def test_control_entire_hierarchy_can_be_renumbered():
    docs = [document(control(), "before"), document(renumber(control(), 20), "after")]
    result = audit(extract(docs))
    assert [f.type for f in result.findings] == ["PRESERVED"]
    assert len(result.unit_changes) == 2
    assert all(u.status == "RETAINED" for u in result.unit_changes)


@pytest.mark.parametrize("mode", ["reported-two-clauses", "all-numbered-headings"])
@pytest.mark.source_documents
def test_real_pair_renumbering_preserves_evidence_outcomes(
    real_documents, real_audit, mode
):
    docs = []
    for original, offset in zip(real_documents, (20, 40)):

        def transform(text):
            if mode == "all-numbered-headings":
                return renumber([text], offset)[0]
            else:
                old, new = (
                    ("5.4.4.", "5.4.14.")
                    if original.version == "before"
                    else ("5.3.3.", "5.3.14.")
                )
                return new + text[len(old) :] if text.startswith(old) else text

        docs.append(rewrite_original(original, transform))
    result = audit(extract(docs))
    assert result.id != real_audit.id
    assert signature(result) == signature(real_audit)
    assert len([f for f in result.findings if f.type == "MOVED"]) == 2
    assert [(u.id, u.status) for u in result.unit_changes] == [
        (u.id, u.status) for u in real_audit.unit_changes
    ]
    spans = {s.id: s for d in result.ir.documents for s in d.spans}
    claims = {c.id: c for c in result.ir.claims}
    for change in result.unit_changes:
        assert {spans[s].version for s in change.source_span_ids} == {"before", "after"}
    for finding in result.findings:
        if finding.type == "MOVED":
            after = claims[finding.after_claim_ids[0]]
            parent = spans[after.source_span_ids[-1]]
            assert "§" + parent.clause in finding.explanation
    validate_record(result)


@pytest.mark.source_documents
def test_paraphrased_transfer_remains_for_review(real_documents):
    # Change only one child formulation, preserving the candidate's context.
    before, after = real_documents
    target = next(s for s in after.spans if s.clause == "5.3.3/б")
    text = target.exact_text.replace("выявления рисков", "обнаружения рисков")
    altered = rewrite_original(
        after, lambda value: text if value == target.exact_text else value
    )
    result = audit(extract([before, altered]))
    claim = next(
        c
        for c in result.ir.claims
        if c.version == "before"
        and c.source_span_ids[0]
        == next(s.id for s in before.spans if s.clause == "5.4.4/б")
    )
    findings = [f for f in result.findings if f.before_claim_id == claim.id]
    assert all(f.type == "REQUIRES_HUMAN_REVIEW" for f in findings)
    search = next(s for s in result.search_results if s.before_claim == claim.id)
    assert any(
        c.after_claim == altered.id + ":" + target.paragraph_id + ":claim"
        and c.relation == "LEXICAL"
        for c in search.candidates
    )


@pytest.mark.parametrize(
    "mutation", ["negated-context", "competing-owner", "wrong-owner-context"]
)
@pytest.mark.source_documents
def test_transfer_needs_full_unambiguous_owner_chain(real_documents, mutation):
    before, after = real_documents
    target = next(s for s in after.spans if s.clause == "5.3.3/б")
    new_owner = "Директоры департаментов и Директоры направлений ДИТААД и ДОА:"
    dnm_owner = (
        "Директор департамента непрерывного мониторинга системы внутреннего контроля:"
    )

    def transform(text):
        if mutation == "negated-context" and text.startswith("5.3.3."):
            return text.replace("взаимодействуют", "не взаимодействуют")
        if mutation == "competing-owner" and text.startswith("5.4.1."):
            return "5.4.1. " + target.exact_text[3:]
        if mutation == "wrong-owner-context":
            if text == "5.3. " + new_owner:
                return "5.3. " + dnm_owner
            if text == "5.4. " + dnm_owner:
                return "5.4. " + new_owner
        return text

    result = audit(extract([before, rewrite_original(after, transform)]))
    source = next(s for s in before.spans if s.clause == "5.4.4/б")
    findings = [f for f in result.findings if f.before_claim_id == source.id + ":claim"]
    assert findings and all(f.type == "REQUIRES_HUMAN_REVIEW" for f in findings)
    validate_record(result)


@pytest.mark.source_documents
def test_retained_old_role_blocks_claim_of_replacement(real_documents):
    before, after = real_documents
    altered = rewrite_original(
        after,
        lambda text: (
            "5.4. Директор направления внутреннего аудита:"
            if text.startswith("5.4. Директор департамента непрерывного мониторинга")
            else text
        ),
    )
    result = audit(extract([before, altered]))
    assert not any(c.status == "REORGANIZED" for c in result.unit_changes)
