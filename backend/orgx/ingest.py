"""Keep original text. Paragraph IDs are physical zero-based OOXML positions.

No OCR, embedded attachments, headers or external content are silently inferred.
"""

import hashlib
import io
import re
import zipfile
from pathlib import Path
from docx import Document as WordDocument
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from .models import Document, SourceSpan, Version

MAX_BYTES = 20 * 1024 * 1024
MAX_EXPANDED = 100 * 1024 * 1024
MAX_SPANS = 12000
NUMBER = re.compile(r"^\s*(\d+(?:\.\d+)*)(?:\.|\s)(.*)", re.S)
LETTER = re.compile(r"^\s*([а-яa-z])[.)]\s+", re.I)


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def strip_label(text: str) -> str:
    return re.sub(
        r"^\s*(?:\d+(?:\.\d+)*\.?|[а-яa-z][.)])\s*", "", text, count=1, flags=re.I
    ).strip()


def normalize_content(text: str) -> str:
    """Normalize typography without stripping substantive leading quantities."""
    return re.sub(r"\s+", " ", text).strip(" ;:.").casefold()


def normalized(text: str) -> str:
    """Normalize a source span with a clause label; use once, before extraction."""
    return normalize_content(strip_label(text))


def _archive_guard(data: bytes):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if sum(x.file_size for x in z.infolist()) > MAX_EXPANDED:
            raise ValueError("Распакованный документ превышает 100 МБ.")
        if len(z.infolist()) > 10000:
            raise ValueError("Слишком много частей в документе.")


def ingest(data: bytes, filename: str, version: Version) -> Document:
    if not data or len(data) > MAX_BYTES:
        raise ValueError("Нужен непустой документ размером до 20 МБ.")
    suffix = Path(filename).suffix.lower()
    sha = digest(data)
    doc_id = f"{version}-{sha[:16]}"
    rows, warnings = [], []
    try:
        if suffix == ".docx":
            _archive_guard(data)
            doc = WordDocument(io.BytesIO(data))
            pnum = tnum = 0
            for item in doc.element.body:
                if item.tag == qn("w:p"):
                    p = Paragraph(item, doc)
                    rows.append(
                        (
                            f"p{pnum:04}",
                            p.text,
                            p.style.name if p.style else "",
                            f"Абзац {pnum + 1}",
                        )
                    )
                    pnum += 1
                elif item.tag == qn("w:tbl"):
                    table = Table(item, doc)
                    seen = set()
                    for ri, row in enumerate(table.rows):
                        for ci, cell in enumerate(row.cells):
                            if cell._tc in seen:
                                continue
                            seen.add(cell._tc)
                            for pi, p in enumerate(cell.paragraphs):
                                rows.append(
                                    (
                                        f"t{tnum}r{ri}c{ci}p{pi}",
                                        p.text,
                                        p.style.name if p.style else "",
                                        f"Таблица {tnum+1}, строка {ri+1}, ячейка {ci+1}",
                                    )
                                )
                    tnum += 1
            if any(
                "embeddings/" in n for n in zipfile.ZipFile(io.BytesIO(data)).namelist()
            ):
                warnings.append(
                    "Вложенные файлы не извлечены; ссылки на приложения не доказывают их содержание."
                )
            if doc.element.xpath(".//w:ins | .//w:del"):
                warnings.append(
                    "Есть исправления Word: требуется принять/отклонить исправления и повторить анализ."
                )
        elif suffix == ".pdf":
            import fitz

            with fitz.open(stream=data, filetype="pdf") as doc:
                for pi, page in enumerate(doc):
                    blocks = sorted(
                        page.get_text("blocks"), key=lambda x: (round(x[1]), x[0])
                    )
                    for bi, block in enumerate(blocks):
                        if block[6] == 0:
                            rows.append(
                                (
                                    f"page{pi+1}b{bi}",
                                    block[4],
                                    "",
                                    f"Страница {pi+1}, блок {bi+1}",
                                )
                            )
            warnings.append(
                "PDF: порядок блоков может отличаться от порядка чтения; OCR не выполняется."
            )
        elif suffix == ".xlsx":
            _archive_guard(data)
            from openpyxl import load_workbook

            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            for si, sheet in enumerate(book):
                for row in sheet:
                    for cell in row:
                        if cell.value is not None:
                            rows.append(
                                (
                                    f"s{si}-{cell.coordinate}",
                                    str(cell.value),
                                    "",
                                    f"{sheet.title}!{cell.coordinate}",
                                )
                            )
            book.close()
            warnings.append(
                "Excel: формулы сохранены как текст; объединённые логические записи требуют проверки."
            )
        else:
            raise ValueError("Поддерживаются DOCX, PDF с текстом и XLSX.")
    except (ValueError,):
        raise
    except Exception as exc:
        raise ValueError(
            "Не удалось прочитать документ. Проверьте формат и целостность файла."
        ) from exc
    if len(rows) > MAX_SPANS:
        raise ValueError("Документ превышает лимит 12 000 фрагментов.")
    spans, section, clause = [], "", ""
    for pid, text, style, locator in rows:
        if not text.strip():
            continue
        number = NUMBER.match(text)
        if number:
            clause = number[1]
            section = clause.split(".")[0]
        letter = LETTER.match(text)
        subclause = f"{clause}/{letter[1]}" if letter else clause
        spans.append(
            SourceSpan(
                id=f"{doc_id}:{pid}",
                document_id=doc_id,
                version=version,
                paragraph_id=pid,
                section=section,
                clause=subclause,
                exact_text=text,
                locator=locator,
                style=style,
            )
        )
    if not spans:
        raise ValueError("Текст не извлечён. Для сканированного PDF сначала нужен OCR.")
    return Document(
        id=doc_id,
        version=version,
        name=Path(filename).name,
        sha256=sha,
        format=suffix[1:],
        spans=spans,
        warnings=warnings,
    )
